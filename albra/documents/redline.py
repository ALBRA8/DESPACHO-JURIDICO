"""
REDLINE DOCX — DOCX con tracked changes nativas (w:ins / w:del).

Sin dependencias pesadas: usamos python-docx (ligero, puro Python).
Para tracked changes nativas usamos elementos XML w:ins y w:del
directamente (python-docx no los expone en su API pública, pero
permite acceso al lxml tree subyacente).

Pipeline V2.0:
1. markdown_draft → docx (placeholder: solo párrafos, sin estilos complejos)
2. add_tracked_insertion(doc, paragraph_idx, text) — añade w:ins
3. add_tracked_deletion(doc, paragraph_idx, text) — añade w:del
4. compare_docx_blackline(doc1_path, doc2_path) → DiffResult
   - Compara paragraph por paragraph
   - Genera un DOCX nuevo con w:ins/w:del que muestra las diferencias
   - Sin dependencia de diff_match_patch ni similar: diff simple por párrafo

Inspiración clean-room (no copiar código): patrón de GLAW y Lawgent.
"""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from albra.observability.logger import get_logger


def _utcnow_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# =========================================================================
# Utilidades XML para construir w:ins / w:del nativos
# =========================================================================

_XML_NS = "http://www.w3.org/XML/1998/namespace"


def _make_ins_run(text: str, author: str, date: str, ins_id: int) -> OxmlElement:
    """Crea un elemento <w:ins> con un <w:r><w:t>text</w:t></w:r> adentro.

    El <w:ins> marca una inserción tracked por Word (visible en
    Track Changes mode).
    """
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), str(ins_id))
    ins.set(qn("w:author"), author)
    ins.set(qn("w:date"), date)
    run = OxmlElement("w:r")
    text_elem = OxmlElement("w:t")
    text_elem.set(f"{{{_XML_NS}}}space", "preserve")
    text_elem.text = text
    run.append(text_elem)
    ins.append(run)
    return ins


def _make_del_run(text: str, author: str, date: str, del_id: int) -> OxmlElement:
    """Crea un elemento <w:del> con un <w:r><w:delText>text</w:delText></w:r>.

    El <w:del> marca una eliminación tracked por Word. El contenido
    usa <w:delText> en vez de <w:t> (spec OOXML).
    """
    dele = OxmlElement("w:del")
    dele.set(qn("w:id"), str(del_id))
    dele.set(qn("w:author"), author)
    dele.set(qn("w:date"), date)
    run = OxmlElement("w:r")
    del_text = OxmlElement("w:delText")
    del_text.set(f"{{{_XML_NS}}}space", "preserve")
    del_text.text = text
    run.append(del_text)
    dele.append(run)
    return dele


def _make_paragraph() -> OxmlElement:
    """Crea un <w:p> vacío."""
    return OxmlElement("w:p")


def _make_paragraph_with_ins(
    text: str, author: str, date: str, ins_id: int
) -> OxmlElement:
    """Crea un <w:p> con un <w:ins> adentro (párrafo nuevo tracked)."""
    p = _make_paragraph()
    p.append(_make_ins_run(text, author, date, ins_id))
    return p


def _make_paragraph_with_del(
    text: str, author: str, date: str, del_id: int
) -> OxmlElement:
    """Crea un <w:p> con un <w:del> adentro."""
    p = _make_paragraph()
    p.append(_make_del_run(text, author, date, del_id))
    return p


def _make_paragraph_normal(text: str) -> OxmlElement:
    """Crea un <w:p> con un <w:r><w:t>text</w:t></w:r> (sin tracked)."""
    p = _make_paragraph()
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(f"{{{_XML_NS}}}space", "preserve")
    t.text = text
    run.append(t)
    p.append(run)
    return p


def _extract_paragraph_text(p_elem: OxmlElement) -> str:
    """Concatena todos los <w:t> del párrafo (texto visible)."""
    parts: list[str] = []
    for t in p_elem.iter(qn("w:t")):
        if t.text:
            parts.append(t.text)
    return "".join(parts)


def _count_elements(root: OxmlElement, tag: str) -> int:
    """Cuenta descendientes con el tag dado (excluye el root)."""
    return sum(1 for el in root.iter(qn(tag)) if el is not root)


# =========================================================================
# RedlineDocxGenerator
# =========================================================================

# Patrones markdown
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\d+\.\s+(.*)$")
_QUOTE_RE = re.compile(r"^>\s?(.*)$")


class RedlineDocxGenerator:
    """Genera DOCX con tracked changes nativas (w:ins / w:del)."""

    _HEADING_STYLES: dict[int, str] = {
        1: "Heading 1",
        2: "Heading 2",
        3: "Heading 3",
        4: "Heading 4",
        5: "Heading 5",
        6: "Heading 6",
    }

    def __init__(self, author: str = "ALBRA-V2", date: str | None = None) -> None:
        """author: nombre que firmará las tracked changes (w:author attribute)."""
        self.author = author
        self.date = date or _utcnow_iso()
        self.log = get_logger("albra.documents.redline")
        # Contador para w:id únicos dentro de cada operación
        self._id_counter = 1000

    def _next_id(self) -> int:
        self._id_counter += 1
        return self._id_counter

    # ===== markdown → docx =====

    def markdown_to_docx(self, markdown_text: str, output_path: str) -> str:
        """Convierte markdown → DOCX. Soporta:
        - # ## ### headings (1-6 '#')
        - párrafos (texto plano)
        - listas con - o * (bullet)
        - listas numeradas con 1. 2.
        - bloques de cita con >

        Returns: output_path (string).
        """
        doc = Document()
        # python-docx crea un párrafo vacío por defecto; lo removemos si
        # el markdown tiene al menos un bloque.
        default_paragraph = doc.paragraphs[0] if doc.paragraphs else None

        added_any = False
        for raw_line in markdown_text.splitlines():
            line = raw_line.rstrip()
            if not line.strip():
                # Línea en blanco: separador (skip)
                continue
            m = _HEADING_RE.match(line)
            if m:
                level = len(m.group(1))
                text = m.group(2).strip()
                style = self._HEADING_STYLES.get(level, "Heading 6")
                doc.add_paragraph(text, style=style)
                added_any = True
                continue
            m = _BULLET_RE.match(line)
            if m:
                text = m.group(1).strip()
                doc.add_paragraph(text, style="List Bullet")
                added_any = True
                continue
            m = _NUMBERED_RE.match(line)
            if m:
                text = m.group(1).strip()
                doc.add_paragraph(text, style="List Number")
                added_any = True
                continue
            m = _QUOTE_RE.match(line)
            if m:
                text = m.group(1).strip()
                # Estilo "Intense Quote" si existe; fallback a "Quote"
                try:
                    doc.add_paragraph(text, style="Intense Quote")
                except KeyError:
                    try:
                        doc.add_paragraph(text, style="Quote")
                    except KeyError:
                        doc.add_paragraph(text)
                added_any = True
                continue
            # Texto normal
            doc.add_paragraph(line)
            added_any = True

        # Remover el párrafo vacío inicial si añadimos bloques
        if added_any and default_paragraph is not None:
            elem = default_paragraph._element
            elem.getparent().remove(elem)

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        self.log.info(
            "redline.markdown_to_docx",
            output_path=str(out),
            paragraphs=len(doc.paragraphs),
        )
        return str(out)

    # ===== add_tracked_insertion =====

    def add_tracked_insertion(
        self, docx_path: str, paragraph_index: int, text: str
    ) -> str:
        """Inserta un párrafo nuevo con w:ins en el índice dado.
        El DOCX queda marcado como 'tracked change' visible en Word.
        Returns: updated path.
        """
        doc = Document(docx_path)
        body = doc.element.body
        paragraphs = body.findall(qn("w:p"))
        if paragraph_index < 0:
            raise ValueError(f"paragraph_index must be >= 0 (got {paragraph_index})")
        if paragraph_index > len(paragraphs):
            raise ValueError(
                f"paragraph_index {paragraph_index} out of range "
                f"(have {len(paragraphs)} paragraphs)"
            )
        new_p = _make_paragraph_with_ins(
            text, author=self.author, date=self.date, ins_id=self._next_id()
        )
        if paragraph_index == len(paragraphs):
            # Append al final (antes del sectPr si existe)
            sect_pr = body.find(qn("w:sectPr"))
            if sect_pr is not None:
                sect_pr.addprevious(new_p)
            else:
                body.append(new_p)
        else:
            paragraphs[paragraph_index].addprevious(new_p)
        doc.save(docx_path)
        self.log.info(
            "redline.add_tracked_insertion",
            docx_path=docx_path,
            paragraph_index=paragraph_index,
            text_preview=text[:50],
        )
        return docx_path

    # ===== add_tracked_deletion =====

    def add_tracked_deletion(
        self,
        docx_path: str,
        paragraph_index: int,
        replacement_text: str = "",
    ) -> str:
        """Marca el párrafo en paragraph_index como w:del (con replacement opcional).
        Returns: updated path.
        """
        doc = Document(docx_path)
        body = doc.element.body
        paragraphs = body.findall(qn("w:p"))
        if paragraph_index < 0 or paragraph_index >= len(paragraphs):
            raise ValueError(
                f"paragraph_index {paragraph_index} out of range "
                f"(have {len(paragraphs)} paragraphs)"
            )
        target_p = paragraphs[paragraph_index]
        # 1. Extraer texto del párrafo (w:t elements)
        full_text = _extract_paragraph_text(target_p)
        # 2. Limpiar el párrafo (conservar solo w:pPr)
        for child in list(target_p):
            if child.tag == qn("w:pPr"):
                continue
            target_p.remove(child)
        # 3. Añadir <w:del> con el texto original
        if full_text:
            del_elem = _make_del_run(
                full_text,
                author=self.author,
                date=self.date,
                del_id=self._next_id(),
            )
            target_p.append(del_elem)
        # 4. Insertar párrafo de reemplazo (si replacement_text no es vacío)
        if replacement_text:
            repl_p = _make_paragraph_with_ins(
                replacement_text,
                author=self.author,
                date=self.date,
                ins_id=self._next_id(),
            )
            target_p.addnext(repl_p)
        doc.save(docx_path)
        self.log.info(
            "redline.add_tracked_deletion",
            docx_path=docx_path,
            paragraph_index=paragraph_index,
            has_replacement=bool(replacement_text),
        )
        return docx_path

    # ===== add_comment (sidecar JSON) =====

    def add_comment(
        self,
        docx_path: str,
        paragraph_index: int,
        comment: str,
        author: str = "ALBRA",
    ) -> str:
        """Añade un comentario (Word comments) al párrafo indicado.
        Returns: updated path.

        V2.0 placeholder: crea archivo sidecar `<docx>.comments.json` con
        la lista de comentarios. Full Word comments requiere XML complejo
        (comments.xml part, relationships, etc.).
        """
        sidecar = Path(f"{docx_path}.comments.json")
        comments: list[dict[str, Any]]
        if sidecar.exists():
            try:
                loaded = json.loads(sidecar.read_text(encoding="utf-8"))
                if isinstance(loaded, list):
                    comments = loaded
                else:
                    comments = []
            except (json.JSONDecodeError, OSError):
                comments = []
        else:
            comments = []
        comments.append({
            "paragraph_index": paragraph_index,
            "comment": comment,
            "author": author,
            "ts": _utcnow_iso(),
        })
        sidecar.write_text(
            json.dumps(comments, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        self.log.info(
            "redline.add_comment",
            docx_path=docx_path,
            paragraph_index=paragraph_index,
            sidecar=str(sidecar),
            total_comments=len(comments),
        )
        return docx_path

    # ===== get_tracked_changes_summary =====

    def get_tracked_changes_summary(self, docx_path: str) -> dict[str, Any]:
        """Lee el DOCX y devuelve un summary de tracked changes:
        {insertions: int, deletions: int, comments: int, authors: list[str]}.
        """
        doc = Document(docx_path)
        body = doc.element.body
        insertions = _count_elements(body, "w:ins")
        deletions = _count_elements(body, "w:del")
        authors_set: set[str] = set()
        for elem in body.iter(qn("w:ins")):
            a = elem.get(qn("w:author"))
            if a:
                authors_set.add(a)
        for elem in body.iter(qn("w:del")):
            a = elem.get(qn("w:author"))
            if a:
                authors_set.add(a)
        # Comments sidecar
        sidecar = Path(f"{docx_path}.comments.json")
        comments_count = 0
        comment_authors: set[str] = set()
        if sidecar.exists():
            try:
                data = json.loads(sidecar.read_text(encoding="utf-8"))
                if isinstance(data, list):
                    comments_count = len(data)
                    for c in data:
                        if isinstance(c, dict) and c.get("author"):
                            comment_authors.add(c["author"])
            except (json.JSONDecodeError, OSError):
                pass
        authors_set.update(comment_authors)
        summary: dict[str, Any] = {
            "insertions": insertions,
            "deletions": deletions,
            "comments": comments_count,
            "authors": sorted(authors_set),
        }
        self.log.info("redline.tracked_changes_summary", docx_path=docx_path, **summary)
        return summary


# =========================================================================
# compare_docx_blackline — diff paragraph por paragraph
# =========================================================================

def _extract_paragraphs(docx_path: str) -> list[str]:
    """Extrae la lista de textos de párrafo del DOCX (en orden)."""
    doc = Document(docx_path)
    return [p.text for p in doc.paragraphs]


def compare_docx_blackline(
    doc1_path: str,
    doc2_path: str,
    output_path: str | None = None,
) -> dict[str, Any]:
    """Compara dos DOCX párrafo por párrafo y genera un blackline DOCX.

    Algoritmo (simple, sin diff_match_patch):
    1. Extrae párrafos de ambos (lista de strings).
    2. diff por párrafo:
       - Si p1[i] == p2[i]: sin cambios
       - Si p1[i] != p2[i]: marcar p1[i] como deletion y p2[i] como insertion
       - Si len(p1) > len(p2): párrafos extra en p1 → deletions
       - Si len(p2) > len(p1): párrafos extra en p2 → insertions
    3. Genera un nuevo DOCX (output_path o '<doc2>_blackline.docx') con
       w:ins y w:del elements mostrando las diferencias.

    Returns: {insertions: int, deletions: int, equal: int, output_path: str}.
    """
    log = get_logger("albra.documents.redline.compare")
    p1 = _extract_paragraphs(doc1_path)
    p2 = _extract_paragraphs(doc2_path)
    author = "ALBRA-BLACKLINE"
    date = _utcnow_iso()
    id_counter = 5000

    def next_id() -> int:
        nonlocal id_counter
        id_counter += 1
        return id_counter

    out_doc = Document()
    body = out_doc.element.body
    # Remover párrafos vacíos iniciales del template
    for p in body.findall(qn("w:p")):
        body.remove(p)
    sect_pr = body.find(qn("w:sectPr"))

    insertions = 0
    deletions = 0
    equal = 0
    max_len = max(len(p1), len(p2))

    new_paragraphs: list[OxmlElement] = []
    for i in range(max_len):
        t1 = p1[i] if i < len(p1) else None
        t2 = p2[i] if i < len(p2) else None
        if t1 is not None and t2 is not None and t1 == t2:
            # Sin cambios: copiar como párrafo normal
            new_paragraphs.append(_make_paragraph_normal(t1))
            equal += 1
        else:
            # Diferente o solo en uno de los dos
            if t1 is not None:
                new_paragraphs.append(
                    _make_paragraph_with_del(t1, author, date, next_id())
                )
                deletions += 1
            if t2 is not None:
                new_paragraphs.append(
                    _make_paragraph_with_ins(t2, author, date, next_id())
                )
                insertions += 1

    # Insertar antes del sectPr (si existe) o append
    for p in new_paragraphs:
        if sect_pr is not None:
            sect_pr.addprevious(p)
        else:
            body.append(p)

    if output_path is None:
        base = Path(doc2_path)
        output_path = str(base.with_name(base.stem + "_blackline" + base.suffix))
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_doc.save(str(out_path))
    log.info(
        "compare_docx_blackline.complete",
        doc1=doc1_path,
        doc2=doc2_path,
        output_path=str(out_path),
        insertions=insertions,
        deletions=deletions,
        equal=equal,
    )
    return {
        "insertions": insertions,
        "deletions": deletions,
        "equal": equal,
        "output_path": str(out_path),
    }
