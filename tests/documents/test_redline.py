"""Tests del RedlineDocxGenerator y compare_docx_blackline."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from albra.documents.redline import (
    RedlineDocxGenerator,
    compare_docx_blackline,
)


@pytest.fixture
def generator():
    return RedlineDocxGenerator(author="test-author")


@pytest.fixture
def sample_markdown():
    return """# Demanda de Tutela

## Antecedentes

El peticionario es mayor de edad.

## Pretensiones

1. Proteger el derecho a la salud.
2. Ordenar a la EPS la entrega del medicamento.
"""


def _write_docx(text: str, path: str) -> str:
    """Helper: escribe un DOCX minimal con un párrafo por línea no vacía."""
    doc = Document()
    body = doc.element.body
    # Remover párrafo vacío inicial del template
    for p in body.findall(qn("w:p")):
        body.remove(p)
    for line in text.splitlines():
        if line.strip():
            doc.add_paragraph(line)
    doc.save(path)
    return path


# ===== markdown_to_docx =====

def test_markdown_to_docx_creates_valid_file(generator, tmp_path, sample_markdown):
    out = generator.markdown_to_docx(sample_markdown, str(tmp_path / "out.docx"))
    assert Path(out).exists()
    # Verificamos que es un DOCX legible
    doc = Document(out)
    assert len(doc.paragraphs) > 0


def test_markdown_to_docx_preserves_headings(generator, tmp_path, sample_markdown):
    out = generator.markdown_to_docx(sample_markdown, str(tmp_path / "out.docx"))
    doc = Document(out)
    texts_styles = [(p.text, p.style.name) for p in doc.paragraphs]
    styles = {s for _, s in texts_styles}
    assert "Heading 1" in styles
    assert "Heading 2" in styles
    h1_texts = [t for t, s in texts_styles if s == "Heading 1"]
    assert any("Demanda de Tutela" in t for t in h1_texts)


def test_markdown_to_docx_preserves_lists(generator, tmp_path):
    md = """Lista bullet:

- Item uno
- Item dos

Lista numerada:

1. Primero
2. Segundo
"""
    out = generator.markdown_to_docx(md, str(tmp_path / "out.docx"))
    doc = Document(out)
    styles = [p.style.name for p in doc.paragraphs]
    assert "List Bullet" in styles
    assert "List Number" in styles


def test_markdown_to_docx_supports_blockquote(generator, tmp_path):
    md = "> Esto es una cita importante."
    out = generator.markdown_to_docx(md, str(tmp_path / "out.docx"))
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert any("cita importante" in t for t in texts)


# ===== add_tracked_insertion =====

def test_add_tracked_insertion_adds_w_ins_element(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    generator.markdown_to_docx("Párrafo uno.", base_path)
    generator.add_tracked_insertion(base_path, 1, "Párrafo insertado")
    doc = Document(base_path)
    body = doc.element.body
    ins_elements = list(body.iter(qn("w:ins")))
    assert len(ins_elements) >= 1
    ins = ins_elements[0]
    assert ins.get(qn("w:author")) == "test-author"


def test_add_tracked_insertion_visible_in_summary(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    generator.markdown_to_docx("Párrafo uno.", base_path)
    generator.add_tracked_insertion(base_path, 1, "Párrafo insertado")
    summary = generator.get_tracked_changes_summary(base_path)
    assert summary["insertions"] >= 1
    assert "test-author" in summary["authors"]


def test_add_tracked_insertion_out_of_range_raises(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    generator.markdown_to_docx("Uno.", base_path)
    with pytest.raises(ValueError, match="out of range"):
        generator.add_tracked_insertion(base_path, 99, "Texto extra")


# ===== add_tracked_deletion =====

def test_add_tracked_deletion_adds_w_del_element(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    md = """Párrafo a borrar.
Párrafo a conservar."""
    generator.markdown_to_docx(md, base_path)
    generator.add_tracked_deletion(base_path, 0)
    doc = Document(base_path)
    body = doc.element.body
    del_elements = list(body.iter(qn("w:del")))
    assert len(del_elements) >= 1
    dele = del_elements[0]
    assert dele.get(qn("w:author")) == "test-author"
    # El texto del w:delText debe preservar el contenido original
    del_texts = [t.text for t in dele.iter(qn("w:delText"))]
    assert any("Párrafo a borrar" in (t or "") for t in del_texts)


def test_add_tracked_deletion_with_replacement(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    md = "Párrafo original."
    generator.markdown_to_docx(md, base_path)
    generator.add_tracked_deletion(base_path, 0, replacement_text="Párrafo nuevo")
    doc = Document(base_path)
    body = doc.element.body
    assert len(list(body.iter(qn("w:del")))) >= 1
    assert len(list(body.iter(qn("w:ins")))) >= 1


def test_add_tracked_deletion_out_of_range_raises(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    generator.markdown_to_docx("Uno.", base_path)
    with pytest.raises(ValueError, match="out of range"):
        generator.add_tracked_deletion(base_path, 99)


# ===== add_comment =====

def test_add_comment_creates_sidecar_json(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    generator.markdown_to_docx("Párrafo uno.", base_path)
    generator.add_comment(base_path, 0, "Revisar redacción", author="abogado-jr")
    sidecar = Path(f"{base_path}.comments.json")
    assert sidecar.exists()
    data = json.loads(sidecar.read_text(encoding="utf-8"))
    assert isinstance(data, list)
    assert len(data) >= 1
    c = data[0]
    assert c["paragraph_index"] == 0
    assert c["comment"] == "Revisar redacción"
    assert c["author"] == "abogado-jr"


def test_add_comment_visible_in_summary(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    generator.markdown_to_docx("Párrafo uno.", base_path)
    generator.add_comment(base_path, 0, "Hola", author="commenter")
    summary = generator.get_tracked_changes_summary(base_path)
    assert summary["comments"] >= 1
    assert "commenter" in summary["authors"]


# ===== get_tracked_changes_summary =====

def test_get_tracked_changes_summary_counts_correctly(generator, tmp_path):
    base_path = str(tmp_path / "base.docx")
    md = """P1
P2
P3"""
    generator.markdown_to_docx(md, base_path)
    # 2 inserciones + 1 delección + 1 comentario
    generator.add_tracked_insertion(base_path, 1, "INS1")
    generator.add_tracked_insertion(base_path, 2, "INS2")
    generator.add_tracked_deletion(base_path, 0)
    generator.add_comment(base_path, 0, "comment")
    summary = generator.get_tracked_changes_summary(base_path)
    assert summary["insertions"] == 2
    assert summary["deletions"] == 1
    assert summary["comments"] == 1
    assert "test-author" in summary["authors"]


# ===== compare_docx_blackline =====

def test_compare_docx_blackline_identical_docs_no_changes(tmp_path):
    text = "Párrafo uno.\nPárrafo dos."
    p1 = _write_docx(text, str(tmp_path / "doc1.docx"))
    p2 = _write_docx(text, str(tmp_path / "doc2.docx"))
    result = compare_docx_blackline(p1, p2, str(tmp_path / "out.docx"))
    assert result["insertions"] == 0
    assert result["deletions"] == 0
    assert result["equal"] == 2
    assert Path(result["output_path"]).exists()


def test_compare_docx_blackline_added_paragraph_is_insertion(tmp_path):
    p1 = _write_docx("Párrafo uno.", str(tmp_path / "doc1.docx"))
    p2 = _write_docx("Párrafo uno.\nPárrafo dos.", str(tmp_path / "doc2.docx"))
    result = compare_docx_blackline(p1, p2, str(tmp_path / "out.docx"))
    assert result["insertions"] >= 1
    assert result["deletions"] == 0


def test_compare_docx_blackline_removed_paragraph_is_deletion(tmp_path):
    p1 = _write_docx("Párrafo uno.\nPárrafo dos.", str(tmp_path / "doc1.docx"))
    p2 = _write_docx("Párrafo uno.", str(tmp_path / "doc2.docx"))
    result = compare_docx_blackline(p1, p2, str(tmp_path / "out.docx"))
    assert result["deletions"] >= 1
    assert result["insertions"] == 0


def test_compare_docx_blackline_modified_paragraph_is_del_then_ins(tmp_path):
    p1 = _write_docx("Texto original.", str(tmp_path / "doc1.docx"))
    p2 = _write_docx("Texto modificado.", str(tmp_path / "doc2.docx"))
    result = compare_docx_blackline(p1, p2, str(tmp_path / "out.docx"))
    assert result["insertions"] == 1
    assert result["deletions"] == 1
    assert result["equal"] == 0
    # Verificamos que el DOCX resultante tiene ambos w:ins y w:del
    doc = Document(result["output_path"])
    body = doc.element.body
    assert len(list(body.iter(qn("w:ins")))) == 1
    assert len(list(body.iter(qn("w:del")))) == 1


def test_compare_docx_blackline_creates_output_file(tmp_path):
    p1 = _write_docx("Uno.", str(tmp_path / "doc1.docx"))
    p2 = _write_docx("Dos.", str(tmp_path / "doc2.docx"))
    # Sin output_path explícito → debe crear '<doc2>_blackline.docx'
    result = compare_docx_blackline(p1, p2)
    out_path = Path(result["output_path"])
    assert out_path.exists()
    assert out_path.stem == "doc2_blackline"
